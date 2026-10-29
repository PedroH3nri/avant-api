"""Relatórios salvos contra o Postgres: CRUD, visibilidade, modelos, execução e exportação."""

import csv
import io
import uuid
from decimal import Decimal

from avant_core.db.uow import SqlAlchemyUnitOfWork
from avant_core.money.services import store_exchange_rates
from avant_core.shared.clock import SystemClock
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine

from tests.conftest import FINANCE, CapturingMailer
from tests.test_api import PASSWORD, _csrf, _onboard_owner, _owner_invite
from tests.test_reports import YESTERDAY, _metrics
from tests.test_team import _member_client


def _report(**extra: object) -> dict[str, object]:
    """Receita, impressões e eCPM de ontem por site, em tabela, para a Empresa inteira."""
    body: dict[str, object] = {
        "name": "Receita por site",
        "query": {
            "metrics": ["revenue", "impressions", "ecpm"],
            "period": {"start": YESTERDAY.isoformat(), "end": YESTERDAY.isoformat()},
            "dimensions": ["site"],
            "order": [{"field": "revenue"}],
        },
        "chart_type": "table",
        "visibility": "company",
    }
    body.update(extra)
    return body


def test_saved_report_crud_run_and_duplicate(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    me = _onboard_owner(client, engine, mailer)
    _metrics(engine, uuid.UUID(me["active_company_id"]))

    created = client.post("/api/v1/reports", json=_report(), headers=_csrf(client))
    assert created.status_code == 201, created.text
    report = created.json()
    assert (report["name"], report["chart_type"], report["visibility"]) == (
        "Receita por site",
        "table",
        "company",
    )
    assert report["editable"] is True and report["owner_user_id"] == me["user"]["id"]
    assert report["query"]["metrics"] == ["revenue", "impressions", "ecpm"]
    assert report["query"]["period"]["start"] == YESTERDAY.isoformat()
    assert "currency" not in report["query"]

    listed = client.get("/api/v1/reports").json()["items"]
    assert [item["id"] for item in listed] == [report["id"]]
    assert client.get(f"/api/v1/reports/{report['id']}").json()["name"] == "Receita por site"

    changed = client.put(
        f"/api/v1/reports/{report['id']}",
        json=_report(name="Por site", visibility="private", description="Ontem"),
        headers=_csrf(client),
    )
    assert changed.status_code == 200, changed.text
    assert (changed.json()["name"], changed.json()["visibility"]) == ("Por site", "private")
    assert changed.json()["description"] == "Ontem"

    run = client.post(f"/api/v1/reports/{report['id']}/run", headers=_csrf(client))
    assert run.status_code == 200, run.text
    assert [row["metrics"]["revenue"] for row in run.json()["rows"]] == ["24", "8", "4"]
    assert run.json()["currency"] == "USD"

    copy = client.post(f"/api/v1/reports/{report['id']}/duplicate", headers=_csrf(client))
    assert copy.status_code == 201, copy.text
    assert (copy.json()["name"], copy.json()["visibility"]) == ("Por site (cópia)", "private")
    assert copy.json()["id"] != report["id"]

    deleted = client.delete(f"/api/v1/reports/{report['id']}", headers=_csrf(client))
    assert deleted.status_code == 204
    missing = client.get(f"/api/v1/reports/{report['id']}")
    assert missing.status_code == 404 and missing.json()["error"]["code"] == "not_found"


def test_visibility_private_company_and_roles(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    _onboard_owner(client, engine, mailer)
    ana = _member_client(client, mailer, "ana@acai.com")
    private = client.post(
        "/api/v1/reports", json=_report(name="Meu", visibility="private"), headers=_csrf(client)
    ).json()
    company = client.post(
        "/api/v1/reports", json=_report(name="De todos"), headers=_csrf(client)
    ).json()
    finance = client.post(
        "/api/v1/reports",
        json=_report(name="Financeiro", visibility="roles", role_ids=[str(FINANCE)]),
        headers=_csrf(client),
    ).json()
    assert finance["role_ids"] == [str(FINANCE)]

    seen = {item["name"]: item for item in ana.get("/api/v1/reports").json()["items"]}
    assert set(seen) == {"De todos", "Financeiro"}
    assert seen["De todos"]["editable"] is False
    hidden = ana.get(f"/api/v1/reports/{private['id']}")
    assert hidden.status_code == 404
    for response in (
        ana.put(f"/api/v1/reports/{company['id']}", json=_report(), headers=_csrf(ana)),
        ana.delete(f"/api/v1/reports/{company['id']}", headers=_csrf(ana)),
        ana.post("/api/v1/reports", json=_report(), headers=_csrf(ana)),
        ana.post(f"/api/v1/reports/{company['id']}/duplicate", headers=_csrf(ana)),
    ):
        assert response.status_code == 403, response.text
        assert response.json()["error"]["code"] == "permission_denied"
    assert ana.post(f"/api/v1/reports/{company['id']}/run", headers=_csrf(ana)).status_code == 200


def test_saved_reports_need_permission(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    _onboard_owner(client, engine, mailer)
    ana = _member_client(client, mailer, "ana@acai.com")
    report = client.post("/api/v1/reports", json=_report(), headers=_csrf(client)).json()
    ana_id = next(
        member["id"]
        for member in client.get("/api/v1/members").json()["items"]
        if member["user"]["name"] == "Ana"
    )
    role = client.post(
        "/api/v1/roles",
        json={"name": "Sem relatórios", "permissions": ["sites.view"]},
        headers=_csrf(client),
    ).json()
    client.patch(f"/api/v1/members/{ana_id}", json={"role_id": role["id"]}, headers=_csrf(client))
    for response in (
        ana.get("/api/v1/reports"),
        ana.get(f"/api/v1/reports/{report['id']}"),
        ana.post(f"/api/v1/reports/{report['id']}/run", headers=_csrf(ana)),
        ana.get(f"/api/v1/reports/{report['id']}/export"),
        ana.get("/api/v1/reports/templates"),
    ):
        assert response.status_code == 403, response.text


def test_invalid_reports_are_refused(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    _onboard_owner(client, engine, mailer)
    pie = client.post("/api/v1/reports", json=_report(chart_type="pie"), headers=_csrf(client))
    assert pie.status_code == 422 and pie.json()["error"]["code"] == "invalid_input"
    internal = _report()
    internal["query"] = {**internal["query"], "metrics": ["revenue_gross"]}  # type: ignore[dict-item]
    refused = client.post("/api/v1/reports", json=internal, headers=_csrf(client))
    assert refused.status_code == 422, refused.text
    no_roles = client.post(
        "/api/v1/reports", json=_report(visibility="roles"), headers=_csrf(client)
    )
    assert no_roles.status_code == 422 and no_roles.json()["error"]["code"] == "invalid_input"
    assert client.get(f"/api/v1/reports/{uuid.uuid4()}").status_code == 404


def test_templates_list_and_create(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    _onboard_owner(client, engine, mailer)
    templates = client.get("/api/v1/reports/templates")
    assert templates.status_code == 200, templates.text
    by_key = {item["key"]: item for item in templates.json()["items"]}
    assert "revenue_by_site" in by_key and "page_views_by_hour" in by_key
    assert by_key["revenue_by_site"]["name"]["en"] == "Revenue by site"
    assert by_key["ecpm_by_day"]["chart_type"] == "line"
    assert by_key["ecpm_by_day"]["query"]["period"]["preset"] == "last_30_days"

    created = client.post("/api/v1/reports/templates/revenue_by_site", headers=_csrf(client))
    assert created.status_code == 201, created.text
    assert (created.json()["name"], created.json()["visibility"]) == (
        "Receita por site",
        "private",
    )
    unknown = client.post("/api/v1/reports/templates/nada", headers=_csrf(client))
    assert unknown.status_code == 404


def test_export_saved_report_csv_and_xlsx(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    me = _onboard_owner(client, engine, mailer)
    _metrics(engine, uuid.UUID(me["active_company_id"]))
    report = client.post("/api/v1/reports", json=_report(), headers=_csrf(client)).json()

    as_csv = client.get(f"/api/v1/reports/{report['id']}/export?locale=en")
    assert as_csv.status_code == 200, as_csv.text
    assert as_csv.headers["content-type"].startswith("text/csv")
    rows = list(csv.reader(io.StringIO(as_csv.text)))
    assert rows[0] == ["Site", "Revenue", "Impressions", "eCPM"]
    assert rows[1] == ["'=blog.com", "24", "3000", "8"]

    as_xlsx = client.get(f"/api/v1/reports/{report['id']}/export?format=xlsx")
    assert as_xlsx.status_code == 200, as_xlsx.text
    assert "spreadsheetml" in as_xlsx.headers["content-type"]
    sheet = load_workbook(io.BytesIO(as_xlsx.content)).active
    assert sheet is not None
    cells = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert cells[0] == ["Site", "Receita", "Impressões", "eCPM"]
    assert cells[1] == ["=blog.com", 24, 3000, 8]
    assert sheet["A2"].data_type == "s"


def test_custom_metric_in_reports_preview_export_and_in_use(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    me = _onboard_owner(client, engine, mailer)
    _metrics(engine, uuid.UUID(me["active_company_id"]))
    metric = {"name": "Metade", "formula": "revenue / 2", "format": "currency"}

    preview = client.post(
        "/api/v1/custom-metrics/validate",
        json={**metric, "sample": {"revenue": "10"}},
        headers=_csrf(client),
    ).json()
    assert preview["valid"] is True and Decimal(preview["value"]) == 5
    empty = client.post(
        "/api/v1/custom-metrics/validate", json={**metric, "sample": {}}, headers=_csrf(client)
    ).json()
    assert empty["valid"] is True and empty["value"] is None
    without = client.post(
        "/api/v1/custom-metrics/validate", json=metric, headers=_csrf(client)
    ).json()
    assert without["value"] is None

    created = client.post("/api/v1/custom-metrics", json=metric, headers=_csrf(client)).json()
    body = _report()
    body["query"] = {**body["query"], "metrics": ["revenue", "custom.metade"]}  # type: ignore[dict-item]
    report = client.post("/api/v1/reports", json=body, headers=_csrf(client))
    assert report.status_code == 201, report.text

    rows = list(
        csv.reader(io.StringIO(client.get(f"/api/v1/reports/{report.json()['id']}/export").text))
    )
    assert rows[0] == ["Site", "Receita", "Metade"]
    assert rows[1] == ["'=blog.com", "24", "12"]

    in_use = client.delete(f"/api/v1/custom-metrics/{created['id']}", headers=_csrf(client))
    assert in_use.status_code == 409
    assert in_use.json()["error"]["code"] == "custom_metric_in_use"


def _ana_id(owner: TestClient) -> str:
    """O id do vínculo da Ana na Empresa."""
    return next(
        str(member["id"])
        for member in owner.get("/api/v1/members").json()["items"]
        if member["user"]["name"] == "Ana"
    )


def test_private_and_other_company_reports_are_404_on_every_route(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    _onboard_owner(client, engine, mailer)
    ana = _member_client(client, mailer, "ana@acai.com")
    manager = client.post(
        "/api/v1/roles",
        json={"name": "Gestão", "permissions": ["reports.view", "reports.manage"]},
        headers=_csrf(client),
    ).json()
    client.patch(
        f"/api/v1/members/{_ana_id(client)}",
        json={"role_id": manager["id"]},
        headers=_csrf(client),
    )
    private = client.post(
        "/api/v1/reports", json=_report(visibility="private"), headers=_csrf(client)
    ).json()

    other = TestClient(client.app)
    token = _owner_invite(engine, mailer, "dono@outra.com", "Outra")
    other.post(
        f"/api/v1/auth/invitations/{token}/accept", json={"name": "Outro", "password": PASSWORD}
    )
    shared = client.post("/api/v1/reports", json=_report(), headers=_csrf(client)).json()

    for person, report_id in ((ana, private["id"]), (other, shared["id"])):
        for response in (
            person.get(f"/api/v1/reports/{report_id}"),
            person.put(f"/api/v1/reports/{report_id}", json=_report(), headers=_csrf(person)),
            person.delete(f"/api/v1/reports/{report_id}", headers=_csrf(person)),
            person.post(f"/api/v1/reports/{report_id}/run", headers=_csrf(person)),
            person.get(f"/api/v1/reports/{report_id}/export"),
            person.post(f"/api/v1/reports/{report_id}/duplicate", headers=_csrf(person)),
        ):
            assert response.status_code == 404, response.text
    assert other.get("/api/v1/reports").json()["items"] == []


def test_manager_edits_and_runs_with_own_sites_and_currency(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    me = _onboard_owner(client, engine, mailer)
    ids = _metrics(engine, uuid.UUID(me["active_company_id"]))
    ana = _member_client(client, mailer, "ana@acai.com")
    manager = client.post(
        "/api/v1/roles",
        json={"name": "Gestão", "permissions": ["reports.view", "reports.manage"]},
        headers=_csrf(client),
    ).json()
    ana_id = _ana_id(client)
    client.patch(
        f"/api/v1/members/{ana_id}", json={"role_id": manager["id"]}, headers=_csrf(client)
    )
    client.put(
        f"/api/v1/members/{ana_id}/sites",
        json={"site_ids": [str(ids["acai"])]},
        headers=_csrf(client),
    )
    report = client.post("/api/v1/reports", json=_report(), headers=_csrf(client)).json()

    seen = ana.get(f"/api/v1/reports/{report['id']}").json()
    assert seen["editable"] is True
    renamed = ana.put(
        f"/api/v1/reports/{report['id']}", json=_report(name="Da Ana"), headers=_csrf(ana)
    )
    assert renamed.status_code == 200 and renamed.json()["owner_user_id"] == me["user"]["id"]

    store_exchange_rates(
        SqlAlchemyUnitOfWork(engine, bypass_rls=True),
        day=YESTERDAY,
        rates={"BRL": Decimal("5")},
        clock=SystemClock(),
    )
    ana.put("/api/v1/auth/preferences", json={"display_currency": "brl"}, headers=_csrf(ana))
    run = ana.post(f"/api/v1/reports/{report['id']}/run", headers=_csrf(ana)).json()
    assert run["currency"] == "BRL"
    assert [row["metrics"]["revenue"] for row in run["rows"]] == ["40"]
    rows = list(csv.reader(io.StringIO(ana.get(f"/api/v1/reports/{report['id']}/export").text)))
    assert [row[0] for row in rows[1:]] == ["acai.com.br"]
    assert (
        client.post(f"/api/v1/reports/{report['id']}/run", headers=_csrf(client)).json()["currency"]
        == "USD"
    )


def test_run_and_export_share_the_query_rate_limit(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    _onboard_owner(client, engine, mailer)
    report = client.post("/api/v1/reports", json=_report(), headers=_csrf(client)).json()
    statuses = [
        client.post(f"/api/v1/reports/{report['id']}/run", headers=_csrf(client)).status_code
        for _ in range(10)
    ]
    exported = client.get(f"/api/v1/reports/{report['id']}/export")
    assert statuses == [200] * 10
    assert exported.status_code == 429 and exported.json()["error"]["code"] == "rate_limited"


def test_chart_options_keep_metric_styles(
    client: TestClient, engine: Engine, mailer: CapturingMailer
) -> None:
    _onboard_owner(client, engine, mailer)
    styles = {
        "metrics": {
            "revenue": {"color": "magenta", "scale": {"mode": "red_green", "from": 1, "to": 10}}
        }
    }
    created = client.post(
        "/api/v1/reports", json=_report(chart_options=styles), headers=_csrf(client)
    )
    assert created.status_code == 201, created.text
    assert client.get(f"/api/v1/reports/{created.json()['id']}").json()["chart_options"] == styles
    pink = {"metrics": {"revenue": {"color": "pink"}}}
    refused = client.post(
        "/api/v1/reports", json=_report(chart_options=pink), headers=_csrf(client)
    )
    assert refused.status_code == 422 and refused.json()["error"]["code"] == "invalid_input"
